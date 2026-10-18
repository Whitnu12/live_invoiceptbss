"""Shared configuration and bootstrap for the central document archive."""

import os
import hashlib
import mimetypes


ARCHIVE_CATEGORIES = {
    "purchase_order": "PO (Purchase Order) Customer",
    "proforma_invoice": "Proforma Invoice (PI)",
    "ctr": "CTR",
    "work_order": "WO (Work Order)",
    "wan": "WAN (Work Authorization Notice)",
    "timesheet": "Timesheet / Daily Report",
    "delivery_note": "Surat Jalan / Manifest",
    "field_report": "Berita Acara Lapangan / Checklist",
    "correspondence": "Korespondensi / Lainnya",
    "official_invoice": "Dokumen Invoice Resmi",
    "tax_invoice": "Dokumen Faktur Pajak",
    "vat_payment": "Dokumen SSP PPN",
    "withholding_tax": "Dokumen Bukti Potong PPh",
}

PRIVATE_CATEGORIES = {
    "official_invoice",
    "tax_invoice",
    "vat_payment",
    "withholding_tax",
}

ARCHIVE_ACCESS_ROLES = {
    "management", "direksi", "finance & tax", "finance", "tax",
    "project manager", "pm", "project support", "super admin", "admin",
}

PRIVATE_ARCHIVE_ROLES = {
    "management", "direksi", "finance & tax", "finance", "tax",
    "project manager", "pm", "super admin", "admin",
}

ALLOWED_ARCHIVE_EXTENSIONS = {
    "pdf", "png", "jpg", "jpeg", "doc", "docx", "xls", "xlsx", "csv", "zip",
}

MAX_ARCHIVE_FILE_BYTES = 25 * 1024 * 1024


def archive_storage_root():
    app_root = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
    return os.path.join(app_root, "database_penyimpanan_aman", "arsip_dokumen_customer")


def bootstrap_document_archive():
    """Create archive metadata and audit tables when the app starts."""
    from db_connection import get_db_connection

    statements = [
        """
        CREATE TABLE IF NOT EXISTS database_document_archive (
            id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
            contract_number VARCHAR(150) NOT NULL,
            pi_number VARCHAR(150) NULL,
            reference_number VARCHAR(150) NULL,
            category VARCHAR(50) NOT NULL,
            original_filename VARCHAR(255) NOT NULL,
            stored_filename VARCHAR(100) NOT NULL,
            storage_path VARCHAR(600) NOT NULL,
            mime_type VARCHAR(150) NULL,
            file_size BIGINT UNSIGNED NOT NULL DEFAULT 0,
            note TEXT NULL,
            uploaded_by VARCHAR(150) NULL,
            uploaded_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
            status VARCHAR(20) NOT NULL DEFAULT 'ACTIVE',
            deleted_by VARCHAR(150) NULL,
            deleted_at DATETIME NULL,
            legacy_key CHAR(64) NULL,
            PRIMARY KEY (id),
            UNIQUE KEY uq_document_archive_stored_filename (stored_filename),
            UNIQUE KEY uq_document_archive_legacy_key (legacy_key),
            KEY idx_document_archive_contract (contract_number),
            KEY idx_document_archive_category (category),
            KEY idx_document_archive_status_uploaded (status, uploaded_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """,
        """
        CREATE TABLE IF NOT EXISTS database_document_archive_audit (
            id BIGINT UNSIGNED NOT NULL AUTO_INCREMENT,
            archive_id BIGINT UNSIGNED NOT NULL,
            actor_username VARCHAR(150) NULL,
            actor_role VARCHAR(80) NULL,
            action VARCHAR(30) NOT NULL,
            detail JSON NULL,
            created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (id),
            KEY idx_document_archive_audit_archive (archive_id, created_at),
            KEY idx_document_archive_audit_actor (actor_username)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """,
        """
        CREATE TABLE IF NOT EXISTS database_document_archive_migration (
            migration_key VARCHAR(100) NOT NULL,
            completed_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (migration_key)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
        """,
    ]
    with get_db_connection() as connection:
        cursor = connection.cursor()
        try:
            for statement in statements:
                cursor.execute(statement)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            cursor.close()
    _import_legacy_archive()


def _import_legacy_archive():
    """Index workbook-backed archive files without moving or duplicating them.

    This scan runs at each application bootstrap. The legacy key makes it
    safe to pick up files added after the first migration without duplicating
    records already present in the central archive.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        return

    archive_root = archive_storage_root()
    workbook_path = os.path.join(archive_root, "metadata_arsip.xlsx")
    if not os.path.isfile(workbook_path):
        return

    from db_connection import get_db_connection

    folder_categories = {
        "01_purchase_order": "purchase_order",
        "02_proforma_invoice": "proforma_invoice",
        "03_ctr": "ctr",
        "04_wo": "work_order",
        "05_wan_sa": "wan",
        "06_timesheet": "timesheet",
        "07_surat_jalan": "delivery_note",
        "08_berita_acara": "field_report",
        "09_korespondensi": "correspondence",
        "10_privat_invoice_resmi": "official_invoice",
        "11_privat_faktur_pajak": "tax_invoice",
        "12_privat_ssp_ppn": "vat_payment",
        "13_privat_bukti_potong_pph": "withholding_tax",
        "02_wan_sa": "wan",
        "03_timesheet": "timesheet",
        "04_surat_jalan": "delivery_note",
        "05_berita_acara": "field_report",
        "06_lainnya": "correspondence",
    }
    label_categories = {label.strip().lower(): key for key, label in ARCHIVE_CATEGORIES.items()}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    rows_to_import = []
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, ())]
        header_indexes = {header.lower(): index for index, header in enumerate(headers) if header}

        def get_value(values, *names):
            for name in names:
                index = header_indexes.get(name.lower())
                if index is not None and index < len(values):
                    value = values[index]
                    if value is not None:
                        return str(value).strip()
            return ""

        for values in iterator:
            raw_path = get_value(values, "Path File")
            if not raw_path:
                continue
            app_root = os.path.dirname(os.path.dirname(archive_root))
            normalized_path = raw_path.replace("\\", "/")
            marker = "database_penyimpanan_aman/"
            marker_index = normalized_path.lower().find(marker)
            if os.path.isabs(raw_path):
                candidate = raw_path
            elif marker_index >= 0:
                candidate = os.path.join(app_root, normalized_path[marker_index:].replace("/", os.sep))
            else:
                candidate = os.path.join(app_root, normalized_path.replace("/", os.sep))
            resolved_path = os.path.realpath(candidate)
            try:
                if os.path.commonpath([os.path.realpath(archive_root), resolved_path]) != os.path.realpath(archive_root):
                    continue
            except ValueError:
                continue
            if not os.path.isfile(resolved_path):
                continue
            folder_name = os.path.basename(os.path.dirname(resolved_path)).strip().lower()
            category = folder_categories.get(folder_name)
            raw_category = get_value(values, "Kategori Dokumen").strip().lower()
            category = category or label_categories.get(raw_category)
            if not category:
                continue
            contract = get_value(values, "Nomor Kontrak")
            if not contract or contract.lower() in {"nan", "-"}:
                continue
            original_name = get_value(values, "Nama File Asli") or os.path.basename(resolved_path)
            try:
                file_size = os.path.getsize(resolved_path)
            except OSError:
                file_size = 0
            uploaded_at = get_value(values, "Tanggal Upload")
            if uploaded_at.lower() in {"", "-", "nan"}:
                uploaded_at = None
            rows_to_import.append((
                contract[:150],
                get_value(values, "Nomor PI", "Nomor PI / PO / Ref")[:150] or None,
                get_value(values, "Nomor PO / Ref", "Nomor PI / PO / Ref")[:150] or None,
                category,
                original_name[:255],
                "legacy-{}".format(hashlib.sha256(resolved_path.encode("utf-8")).hexdigest()[:36]),
                resolved_path,
                mimetypes.guess_type(original_name)[0],
                file_size,
                get_value(values, "Keterangan") or None,
                uploaded_at,
                hashlib.sha256(resolved_path.encode("utf-8")).hexdigest(),
            ))
    finally:
        workbook.close()

    with get_db_connection() as connection:
        cursor = connection.cursor()
        try:
            if rows_to_import:
                cursor.executemany(
                    """
                    INSERT IGNORE INTO database_document_archive
                    (contract_number, pi_number, reference_number, category, original_filename,
                     stored_filename, storage_path, mime_type, file_size, note, uploaded_at, legacy_key)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, COALESCE(%s, CURRENT_TIMESTAMP), %s)
                    """,
                    rows_to_import,
                )
            cursor.execute(
                "INSERT IGNORE INTO database_document_archive_migration (migration_key) VALUES (%s)",
                ("legacy_excel_archive_import_v1",),
            )
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            cursor.close()
