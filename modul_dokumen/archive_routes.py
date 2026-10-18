"""Routes for the central customer and supporting document archive."""

import csv
import io
import json
import logging
import mimetypes
import os
import uuid
import zipfile
from urllib.parse import quote
from xml.etree import ElementTree

from flask import Blueprint, jsonify, make_response, render_template, request, send_file, session
from werkzeug.utils import secure_filename

from db_connection import get_db_connection
from modul_dokumen.archive_service import (
    ALLOWED_ARCHIVE_EXTENSIONS,
    ARCHIVE_ACCESS_ROLES,
    ARCHIVE_CATEGORIES,
    MAX_ARCHIVE_FILE_BYTES,
    PRIVATE_ARCHIVE_ROLES,
    PRIVATE_CATEGORIES,
    archive_storage_root,
)
from modul_keamanan.decorators import login_required
from modul_input.modul_1_database import can_write


archive_bp = Blueprint("archive", __name__)
LOGGER = logging.getLogger(__name__)


def _role():
    return str(session.get("current_role", "")).strip().lower()


def _has_archive_access():
    return _role() in ARCHIVE_ACCESS_ROLES


def _file_path(storage_path):
    root = os.path.realpath(archive_storage_root())
    path = os.path.realpath(str(storage_path or ""))
    try:
        return path if os.path.commonpath([root, path]) == root else None
    except ValueError:
        return None


def _audit(archive_id, action, detail=None):
    try:
        with get_db_connection() as connection:
            cursor = connection.cursor()
            try:
                cursor.execute(
                    """
                    INSERT INTO database_document_archive_audit
                    (archive_id, actor_username, actor_role, action, detail)
                    VALUES (%s, %s, %s, %s, %s)
                    """,
                    (
                        archive_id,
                        str(session.get("current_user", "") or "")[:150] or None,
                        str(session.get("current_role", "") or "")[:80] or None,
                        action,
                        json.dumps(detail or {}, ensure_ascii=False),
                    ),
                )
                connection.commit()
            finally:
                cursor.close()
    except Exception:
        LOGGER.warning("Audit arsip dokumen gagal dicatat: %s", action, exc_info=True)


def _forbidden_json():
    return jsonify({"success": False, "message": "Akses arsip dokumen ditolak."}), 403


def _preview_content(path, extension):
    """Build bounded HTML-friendly previews without writing converted copies."""
    if extension == "csv":
        with open(path, "rb") as source:
            raw = source.read(5 * 1024 * 1024 + 1)
        if len(raw) > 5 * 1024 * 1024:
            return "message", "Pratinjau CSV dibatasi untuk berkas maksimal 5 MB. Berkas asli tetap tersimpan tanpa perubahan."
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp1252", errors="replace")
        rows = list(csv.reader(io.StringIO(content)))[:101]
        return "tables", [{"title": "Data CSV", "rows": rows[:100], "truncated": len(rows) > 100}]

    if extension == "xlsx":
        with zipfile.ZipFile(path) as package:
            infos = package.infolist()
            if sum(item.file_size for item in infos) > 150 * 1024 * 1024:
                return "message", "Pratinjau Excel dilewati karena ukuran isi workbook terlalu besar. Berkas asli tetap tersimpan tanpa perubahan."
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sections = []
        sheet_count = len(workbook.worksheets)
        try:
            for sheet in workbook.worksheets[:5]:
                rows = []
                column_truncated = False
                for row in sheet.iter_rows(max_row=101, max_col=21, values_only=True):
                    column_truncated = column_truncated or len(row) > 20
                    rows.append(["" if value is None else str(value) for value in row[:20]])
                truncated = len(rows) > 100 or column_truncated
                sections.append({"title": sheet.title, "rows": rows[:100], "truncated": truncated})
        finally:
            workbook.close()
        if sheet_count > 5:
            sections.append({"title": "Catatan", "message": "Pratinjau dibatasi hingga 5 sheet."})
        return "tables", sections

    if extension == "docx":
        with zipfile.ZipFile(path) as package:
            try:
                document = package.getinfo("word/document.xml")
            except KeyError:
                return "message", "Isi dokumen Word tidak dapat dibaca. Berkas aslinya tetap tersedia untuk diunduh."
            if document.file_size > 8 * 1024 * 1024:
                return "message", "Pratinjau Word dilewati karena konten dokumen terlalu besar. Berkas asli tetap tersimpan tanpa perubahan."
            root = ElementTree.fromstring(package.read(document))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        body = root.find("w:body", ns)
        paragraphs, tables = [], []
        body_blocks = list(body) if body is not None else []
        if body is not None:
            for block in body_blocks[:501]:
                if block.tag.endswith("}p"):
                    text = "".join(node.text or "" for node in block.findall(".//w:t", ns)).strip()
                    if text:
                        paragraphs.append(text)
                elif block.tag.endswith("}tbl"):
                    rows = []
                    for row in block.findall("w:tr", ns)[:100]:
                        cells = []
                        for cell in row.findall("w:tc", ns)[:20]:
                            cells.append(" ".join((node.text or "") for node in cell.findall(".//w:t", ns)).strip())
                        rows.append(cells)
                    if rows:
                        tables.append({"title": "Tabel dokumen", "rows": rows, "truncated": len(block.findall("w:tr", ns)) > 100})
        return "docx", {"paragraphs": paragraphs[:500], "tables": tables, "truncated": len(body_blocks) > 500}

    if extension == "zip":
        with zipfile.ZipFile(path) as package:
            files = package.infolist()
        rows = [[item.filename, "Direktori" if item.is_dir() else _human_file_size(item.file_size)] for item in files[:200]]
        return "tables", [{"title": "Isi arsip ZIP", "headers": ["Nama", "Ukuran"], "rows": rows, "truncated": len(files) > 200}]

    return "message", "Pratinjau browser belum tersedia untuk format .{}. Berkas asli tetap tersimpan; gunakan Unduh bila perlu membuka dengan aplikasi lokal.".format(extension or "ini")


def _human_file_size(size):
    size = int(size or 0)
    if size < 1024:
        return "{} B".format(size)
    if size < 1024 * 1024:
        return "{} KB".format(round(size / 1024))
    return "{:.1f} MB".format(size / (1024 * 1024))


@archive_bp.route("/modul/arsip")
@login_required
def index():
    if not _has_archive_access():
        return "Anda tidak memiliki hak akses ke arsip dokumen.", 403
    categories = {
        key: label for key, label in ARCHIVE_CATEGORIES.items()
        if key not in PRIVATE_CATEGORIES or _role() in PRIVATE_ARCHIVE_ROLES
    }
    return render_template(
        "archive_documents.html",
        title="Arsip Dokumen",
        module_code="PUSAT ARSIP",
        categories=categories,
        can_write=can_write(),
    )


@archive_bp.route("/api/archive-documents", methods=["GET"])
@login_required
def list_documents():
    if not _has_archive_access():
        return _forbidden_json()
    try:
        query = str(request.args.get("q", "") or "").strip()[:150]
        category = str(request.args.get("category", "") or "").strip()
        contract = str(request.args.get("contract", "") or "").strip()[:150]
        page = max(1, int(request.args.get("page", "1")))
        page_size = min(100, max(10, int(request.args.get("page_size", "25"))))
        if category and category not in ARCHIVE_CATEGORIES:
            return jsonify({"success": False, "message": "Kategori dokumen tidak valid."}), 400
        if category in PRIVATE_CATEGORIES and _role() not in PRIVATE_ARCHIVE_ROLES:
            return _forbidden_json()

        conditions = ["status = 'ACTIVE'"]
        params = []
        if _role() not in PRIVATE_ARCHIVE_ROLES:
            conditions.append("category NOT IN (%s, %s, %s, %s)")
            params.extend(sorted(PRIVATE_CATEGORIES))
        if category:
            conditions.append("category = %s")
            params.append(category)
        if contract:
            conditions.append("contract_number = %s")
            params.append(contract)
        if query:
            escaped = query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            conditions.append(
                "(contract_number LIKE %s OR pi_number LIKE %s OR reference_number LIKE %s "
                "OR original_filename LIKE %s OR note LIKE %s)"
            )
            params.extend(["%{}%".format(escaped)] * 5)
        where_sql = " AND ".join(conditions)
        offset = (page - 1) * page_size
        with get_db_connection() as connection:
            cursor = connection.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT COUNT(*) AS total FROM database_document_archive WHERE " + where_sql,
                    tuple(params),
                )
                total = int(cursor.fetchone()["total"])
                cursor.execute(
                    """
                    SELECT id, contract_number, pi_number, reference_number, category,
                           original_filename, mime_type, file_size, note, uploaded_by, uploaded_at
                    FROM database_document_archive WHERE """ + where_sql +
                    " ORDER BY uploaded_at DESC, id DESC LIMIT %s OFFSET %s",
                    tuple(params + [page_size, offset]),
                )
                rows = cursor.fetchall()
            finally:
                cursor.close()
        for row in rows:
            row["category_label"] = ARCHIVE_CATEGORIES.get(row.get("category"), row.get("category"))
            if row.get("uploaded_at"):
                row["uploaded_at"] = row["uploaded_at"].isoformat(sep=" ", timespec="minutes")
            row["can_delete"] = can_write()
        return jsonify({
            "success": True,
            "data": rows,
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": (total + page_size - 1) // page_size,
        })
    except (ValueError, TypeError) as error:
        return jsonify({"success": False, "message": str(error)}), 400
    except Exception:
        LOGGER.exception("API daftar arsip dokumen gagal.")
        return jsonify({"success": False, "message": "Gagal memuat daftar arsip dokumen."}), 500


@archive_bp.route("/api/archive-documents", methods=["POST"])
@login_required
def upload_document():
    if not _has_archive_access():
        return _forbidden_json()
    if not can_write():
        return jsonify({"success": False, "message": "Role Anda hanya memiliki akses baca."}), 403
    role = _role()
    file_obj = request.files.get("file")
    contract = str(request.form.get("contract_number", "") or "").strip()[:150]
    pi_number = str(request.form.get("pi_number", "") or "").strip()[:150]
    reference_number = str(request.form.get("reference_number", "") or "").strip()[:150]
    category = str(request.form.get("category", "") or "").strip()
    note = str(request.form.get("note", "") or "").strip()
    if not contract:
        return jsonify({"success": False, "message": "Nomor kontrak wajib diisi."}), 400
    if category not in ARCHIVE_CATEGORIES:
        return jsonify({"success": False, "message": "Kategori dokumen tidak valid."}), 400
    if category in PRIVATE_CATEGORIES and role not in PRIVATE_ARCHIVE_ROLES:
        return _forbidden_json()
    if not file_obj or not file_obj.filename:
        return jsonify({"success": False, "message": "Berkas wajib dipilih."}), 400
    original_name = secure_filename(file_obj.filename)[:255]
    extension = original_name.rsplit(".", 1)[-1].lower() if "." in original_name else ""
    if not original_name or extension not in ALLOWED_ARCHIVE_EXTENSIONS:
        return jsonify({"success": False, "message": "Format file tidak didukung."}), 400
    file_bytes = file_obj.read(MAX_ARCHIVE_FILE_BYTES + 1)
    if not file_bytes:
        return jsonify({"success": False, "message": "Berkas kosong."}), 400
    if len(file_bytes) > MAX_ARCHIVE_FILE_BYTES:
        return jsonify({"success": False, "message": "Ukuran file maksimal 25 MB."}), 400

    stored_name = "{}.{}".format(uuid.uuid4().hex, extension)
    archive_root = archive_storage_root()
    category_dir = os.path.join(archive_root, category)
    os.makedirs(category_dir, exist_ok=True)
    storage_path = os.path.join(category_dir, stored_name)
    try:
        with open(storage_path, "xb") as output:
            output.write(file_bytes)
        with get_db_connection() as connection:
            cursor = connection.cursor()
            try:
                cursor.execute(
                    """
                    INSERT INTO database_document_archive
                    (contract_number, pi_number, reference_number, category, original_filename,
                     stored_filename, storage_path, mime_type, file_size, note, uploaded_by)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        contract, pi_number or None, reference_number or None, category,
                        original_name, stored_name, storage_path,
                        file_obj.mimetype or mimetypes.guess_type(original_name)[0],
                        len(file_bytes), note or None,
                        str(session.get("current_user", "") or "")[:150] or None,
                    ),
                )
                archive_id = cursor.lastrowid
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            finally:
                cursor.close()
        _audit(archive_id, "UPLOAD", {"category": category, "filename": original_name})
        return jsonify({"success": True, "message": "Dokumen berhasil diarsipkan.", "data": {"id": archive_id}}), 201
    except Exception:
        if os.path.exists(storage_path):
            try:
                os.remove(storage_path)
            except OSError:
                LOGGER.exception("Berkas arsip gagal dibersihkan setelah upload gagal.")
        LOGGER.exception("Upload arsip dokumen gagal.")
        return jsonify({"success": False, "message": "Gagal menyimpan dokumen. Silakan coba kembali."}), 500


@archive_bp.route("/api/archive-documents/<int:archive_id>/download", methods=["GET"])
@login_required
def download_document(archive_id):
    if not _has_archive_access():
        return _forbidden_json()
    try:
        with get_db_connection() as connection:
            cursor = connection.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT category, original_filename, storage_path FROM database_document_archive "
                    "WHERE id = %s AND status = 'ACTIVE' LIMIT 1",
                    (archive_id,),
                )
                row = cursor.fetchone()
            finally:
                cursor.close()
        if not row:
            return jsonify({"success": False, "message": "Arsip tidak ditemukan."}), 404
        if row["category"] in PRIVATE_CATEGORIES and _role() not in PRIVATE_ARCHIVE_ROLES:
            return _forbidden_json()
        path = _file_path(row["storage_path"])
        if not path or not os.path.isfile(path):
            return jsonify({"success": False, "message": "File arsip tidak tersedia."}), 404
        _audit(archive_id, "DOWNLOAD", {})
        return send_file(path, as_attachment=True, download_name=row["original_filename"])
    except Exception:
        LOGGER.exception("Download arsip dokumen gagal.")
        return jsonify({"success": False, "message": "Gagal mengunduh file arsip."}), 500


@archive_bp.route("/api/archive-documents/<int:archive_id>/preview", methods=["GET"])
@login_required
def preview_document(archive_id):
    if not _has_archive_access():
        return "Akses arsip dokumen ditolak.", 403
    try:
        with get_db_connection() as connection:
            cursor = connection.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT category, original_filename, storage_path, file_size FROM database_document_archive "
                    "WHERE id = %s AND status = 'ACTIVE' LIMIT 1",
                    (archive_id,),
                )
                row = cursor.fetchone()
            finally:
                cursor.close()
        if not row:
            return "Arsip tidak ditemukan.", 404
        if row["category"] in PRIVATE_CATEGORIES and _role() not in PRIVATE_ARCHIVE_ROLES:
            return "Akses kategori privat ditolak.", 403
        path = _file_path(row["storage_path"])
        if not path or not os.path.isfile(path):
            return "File arsip tidak tersedia.", 404

        original_name = row["original_filename"]
        extension = os.path.splitext(original_name)[1].lower().lstrip(".")
        if extension in {"pdf", "png", "jpg", "jpeg"}:
            content_type = mimetypes.guess_type(original_name)[0] or "application/octet-stream"
            response = send_file(
                path,
                mimetype=content_type,
                as_attachment=False,
                download_name=original_name,
                conditional=True,
                max_age=0,
            )
            response.headers["Content-Disposition"] = "inline; filename*=UTF-8''{}".format(quote(original_name, safe=""))
            response.headers["X-Content-Type-Options"] = "nosniff"
            response.headers["Cache-Control"] = "private, no-store"
            _audit(archive_id, "VIEW", {"filename": original_name})
            return response

        preview_type, preview_data = _preview_content(path, extension)
        response = render_template(
            "archive_preview.html",
            file_name=original_name,
            category_label=ARCHIVE_CATEGORIES.get(row["category"], row["category"]),
            preview_type=preview_type,
            preview_data=preview_data,
            file_size=_human_file_size(row["file_size"]),
        )
        _audit(archive_id, "VIEW", {"filename": original_name})
        html_response = make_response(response)
        html_response.headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; frame-ancestors 'self'; base-uri 'none'"
        html_response.headers["X-Content-Type-Options"] = "nosniff"
        html_response.headers["Cache-Control"] = "private, no-store"
        return html_response
    except Exception:
        LOGGER.exception("Pratinjau arsip dokumen gagal.")
        return "Pratinjau gagal dibuka. Berkas asli tetap berada di arsip.", 500


@archive_bp.route("/api/archive-documents/<int:archive_id>", methods=["DELETE"])
@login_required
def delete_document(archive_id):
    if not _has_archive_access():
        return _forbidden_json()
    if not can_write():
        return jsonify({"success": False, "message": "Role Anda hanya memiliki akses baca."}), 403
    try:
        with get_db_connection() as connection:
            cursor = connection.cursor(dictionary=True)
            try:
                cursor.execute(
                    "SELECT category, original_filename, storage_path FROM database_document_archive "
                    "WHERE id = %s AND status = 'ACTIVE' LIMIT 1",
                    (archive_id,),
                )
                row = cursor.fetchone()
                if not row:
                    return jsonify({"success": False, "message": "Arsip tidak ditemukan."}), 404
                if row["category"] in PRIVATE_CATEGORIES and _role() not in PRIVATE_ARCHIVE_ROLES:
                    return _forbidden_json()
                cursor.execute(
                    "UPDATE database_document_archive SET status='DELETED', deleted_by=%s, "
                    "deleted_at=CURRENT_TIMESTAMP WHERE id=%s AND status='ACTIVE'",
                    (str(session.get("current_user", "") or "")[:150] or None, archive_id),
                )
                if cursor.rowcount == 0:
                    connection.rollback()
                    return jsonify({"success": False, "message": "Arsip sudah dihapus."}), 404
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            finally:
                cursor.close()
        path = _file_path(row["storage_path"])
        if not path:
            raise OSError("Lokasi arsip tidak berada dalam direktori penyimpanan yang diizinkan.")
        if path and os.path.isfile(path):
            os.remove(path)
        _audit(archive_id, "DELETE", {"filename": row["original_filename"]})
        return jsonify({"success": True, "message": "Dokumen dan file fisiknya berhasil dihapus."})
    except OSError:
        LOGGER.exception("File fisik arsip gagal dihapus.")
        _audit(archive_id, "DELETE_FILE_ERROR", {})
        return jsonify({"success": False, "message": "Arsip dinonaktifkan, tetapi file fisik gagal dihapus. Hubungi administrator."}), 500
    except Exception:
        LOGGER.exception("Hapus arsip dokumen gagal.")
        return jsonify({"success": False, "message": "Gagal menghapus arsip dokumen."}), 500
